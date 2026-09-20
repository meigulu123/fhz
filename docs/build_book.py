# -*- coding: utf-8 -*-
"""把 docs/book/ 下的 markdown 书稿生成《知途伴学》项目书 Word。

用法：
    python docs/build_book.py             # 生成（缺图时跳过并告警）
    python docs/build_book.py --strict    # 缺图即报错退出

输出：docs/知途伴学_项目书.docx
说明：目录为 Word TOC 域，打开文档后 Ctrl+A 全选再按 F9 刷新即可生成页码。
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent          # E:\guochuang
BOOK_DIR = ROOT / "docs" / "book"
LOGO_PATH = ROOT / "assets" / "logo.png"
OUT_PATH = ROOT / "docs" / "知途伴学_项目书.docx"

SONG = "宋体"
HEI = "黑体"
LATIN = "Times New Roman"
MONO = "Consolas"
DARK = RGBColor(0x1F, 0x4E, 0x79)

stats = {"paragraphs": 0, "tables": 0, "images": 0, "missing": []}
STRICT = "--strict" in sys.argv


# ---------------------------------------------------------------- 底层工具
def _run_font(run, east=SONG, latin=LATIN, size=12, bold=False, italic=False,
              color=None):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), east)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)


def _style_ea(style, east=SONG, latin=LATIN, size=12, bold=False,
              color=None):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), east)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)


def _indent_chars(p, chars=2):
    """按字符数设置首行缩进（Word 原生 w:firstLineChars）。"""
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars * 100))
    ind.set(qn("w:firstLine"), str(int(chars * 240)))   # 兜底：24pt=480 twips


def _shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
BOLD_FULL = re.compile(r"^\*\*(.+?)\*\*$")


def add_rich(p, text, size=12, east=SONG, bold=False):
    """带 **粗体** 与 `行内代码` 的正文段。"""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = p.add_run(part[2:-2])
            _run_font(run, east=east, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = p.add_run(part[1:-1])
            _run_font(run, east=east, latin=MONO, size=size - 1)
        else:
            run = p.add_run(part)
            _run_font(run, east=east, size=size, bold=bold)


# ---------------------------------------------------------------- 文档骨架
def setup_document():
    doc = Document()
    for sec in doc.sections:
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)   # A4
        sec.top_margin = sec.bottom_margin = Cm(2.3)
        sec.left_margin = sec.right_margin = Cm(2.3)

    normal = doc.styles["Normal"]
    _style_ea(normal, east=SONG, size=12)
    pf = normal.paragraph_format
    pf.line_spacing = 1.25
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    h1 = doc.styles["Heading 1"]          # 章标题：黑体三号居中
    _style_ea(h1, east=HEI, size=16)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(12)
    h1.paragraph_format.line_spacing = 1.3
    h1.paragraph_format.page_break_before = True

    h2 = doc.styles["Heading 2"]          # 节标题：黑体四号左对齐
    _style_ea(h2, east=HEI, size=14)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.line_spacing = 1.3

    h3 = doc.styles["Heading 3"]          # 条标题：黑体小四左对齐
    _style_ea(h3, east=HEI, size=12)
    h3.paragraph_format.space_before = Pt(4)
    h3.paragraph_format.space_after = Pt(2)
    h3.paragraph_format.line_spacing = 1.3
    return doc


def cover(doc):
    def line(text, east, size, before=0, after=0, bold=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        run = p.add_run(text)
        _run_font(run, east=east, size=size, bold=bold, color=DARK
                  if east == HEI else None)

    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO_PATH), height=Cm(3.6))
    line("中国国际大学生创新大赛（2026）", HEI, 22, before=24, after=6)
    line("产业命题赛道 · 项目书", HEI, 15, after=30)
    line("知途伴学", HEI, 30, after=10)
    line("大模型驱动的自适应学习路径决策与伴学智能体开发", HEI, 16,
         after=8)
    line("（命题九十六 · 命题企业：科大讯飞）", HEI, 13, after=60)
    line("团队：广东理工学院 · 知途伴学", SONG, 14, after=10)
    line("创始人：曹锦杰", SONG, 14, after=40)
    line("2026 年 9 月", SONG, 14)


def toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("目  录")
    _run_font(run, east=HEI, size=16)

    p = doc.add_paragraph()
    run = p.add_run()
    for el, attrs, text in (
            ("fldChar", {"w:fldCharType": "begin"}, None),
            ("instrText", {"xml:space": "preserve"},
             'TOC \\o "1-2" \\h \\z \\u'),
            ("fldChar", {"w:fldCharType": "separate"}, None),
            ("t", {}, "打开 Word 后按 Ctrl+A 全选，再按 F9 刷新生成目录"),
            ("fldChar", {"w:fldCharType": "end"}, None)):
        node = OxmlElement("w:" + el)
        for k, v in attrs.items():
            node.set(qn(k), v)
        if text is not None:
            node.text = text
        run._r.append(node)


def footer_pagenum(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    _run_font(run, size=9)
    for el, attrs, text in (
            ("fldChar", {"w:fldCharType": "begin"}, None),
            ("instrText", {"xml:space": "preserve"}, "PAGE"),
            ("fldChar", {"w:fldCharType": "end"}, None)):
        node = OxmlElement("w:" + el)
        for k, v in attrs.items():
            node.set(qn(k), v)
        if text is not None:
            node.text = text
        run._r.append(node)


# ---------------------------------------------------------------- 内容渲染
def render_md(doc, path):
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    i, n = 0, len(lines)

    while i < n:
        line = lines[i]

        # 代码块
        if line.strip().startswith("```"):
            buf = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1          # 跳过结尾 ```
            for code_line in buf:
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.line_spacing = 1.0
                pf.left_indent = Cm(0.4)
                pf.right_indent = Cm(0.4)
                _shade(p, "F5F5F5")
                run = p.add_run(code_line if code_line else " ")
                _run_font(run, east=MONO, latin=MONO, size=9)
                stats["paragraphs"] += 1
            continue

        s = line.strip()

        if not s:
            i += 1
            continue

        if s.startswith("# "):                       # 章标题
            p = doc.add_paragraph(s[2:], style="Heading 1")
            stats["paragraphs"] += 1
        elif s.startswith("## "):                    # 节标题
            p = doc.add_paragraph(s[3:], style="Heading 2")
            stats["paragraphs"] += 1
        elif s.startswith("### "):                   # 条标题
            p = doc.add_paragraph(s[4:], style="Heading 3")
            stats["paragraphs"] += 1
        elif s.startswith("[F] "):                   # 公式行
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(s[4:])
            _run_font(run, east=LATIN, latin=LATIN, size=12, italic=True)
            stats["paragraphs"] += 1
        elif s.startswith("|"):                      # 表格
            i = add_table(doc, lines, i)
            continue
        elif s.startswith("!["):                     # 图片
            add_image(doc, s)
        elif BOLD_FULL.match(s):                     # 独立粗体行 = 表题
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(BOLD_FULL.match(s).group(1))
            _run_font(run, east=HEI, size=10.5)
            stats["paragraphs"] += 1
        elif s.startswith("- "):                     # 列表项
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.first_line_indent = Cm(-0.74)
            add_rich(p, "• " + s[2:], size=12)
            stats["paragraphs"] += 1
        else:                                        # 正文段
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _indent_chars(p, 2)
            add_rich(p, s, size=12)
            stats["paragraphs"] += 1
        i += 1


def add_table(doc, lines, i):
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    if not rows:
        return i
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    cellMar = OxmlElement("w:tblCellMar")
    for side, val in (("top", "0"), ("left", "57"), ("bottom", "0"),
                      ("right", "57")):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), val)
        el.set(qn("w:type"), "dxa")
        cellMar.append(el)
    tblPr.append(cellMar)
    for r_i, row in enumerate(rows):
        cells = table.add_row().cells
        for c_i in range(ncols):
            cell = cells[c_i]
            cell.paragraphs[0].text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if c_i < len(row):
                add_rich(p, row[c_i], size=8.5, bold=(r_i == 0))
            if r_i == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "EAF1F8")
                tcPr.append(shd)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)
    stats["tables"] += 1
    stats["paragraphs"] += len(rows)
    return i


def add_image(doc, s):
    m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", s)
    if not m:
        return 1
    caption, rel = m.group(1), m.group(2)
    img = (BOOK_DIR / rel).resolve()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    if img.exists():
        p.add_run().add_picture(str(img), width=Cm(11.5))
        stats["images"] += 1
    else:
        stats["missing"].append(str(img))
        run = p.add_run(f"[缺图：{img.name}]")
        _run_font(run, east=SONG, size=9, color=RGBColor(0xC0, 0, 0))
        if STRICT:
            sys.exit(f"missing image: {img}")
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(3)
        cp.paragraph_format.space_after = Pt(8)
        run = cp.add_run(caption)
        _run_font(run, east=HEI, size=10.5)
        stats["paragraphs"] += 1
    return 0


def main():
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else OUT_PATH
    doc = setup_document()
    cover(doc)
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    footer_pagenum(body)
    toc(doc)
    for f in sorted(BOOK_DIR.glob("*.md")):
        render_md(doc, f)
    doc.core_properties.title = "知途伴学——大模型驱动的自适应学习路径决策与伴学智能体开发（项目书）"
    doc.core_properties.author = "广东理工学院 · 知途伴学"
    doc.save(out_path)
    print(f"OK -> {out_path}")
    print(f"段落 {stats['paragraphs']} | 表格 {stats['tables']} | 图片 {stats['images']}")
    if stats["missing"]:
        print(f"缺图 {len(stats['missing'])} 张（生成时以占位符代替）：")
        for m in stats["missing"]:
            print("  -", m)


if __name__ == "__main__":
    main()
